"""Contract between each interface and the same released validation package."""
import ast
import asyncio
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from agent_ref_validator import core
from agent_ref_validator.exports import build_excel_workbook
from agent_ref_validator.verify import verify_release


class SharedIntegrationTests(unittest.TestCase):
    def test_release_manifest_matches_the_executed_core(self):
        manifest = verify_release(REPO / "agent_ref_validator")
        self.assertEqual(manifest["version"], core.VALIDATOR_VERSION)
        self.assertEqual(manifest["files"]["core.py"], core.VALIDATOR_SHA256)

    def test_modified_release_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            package = Path(directory)
            (package / "core.py").write_text("changed = True\n", encoding="utf-8")
            (package / "release.json").write_text(json.dumps({"files": {"core.py": "wrong"}}), encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "Validator release drift"):
                verify_release(package)

    def test_interface_uses_shared_validation_functions(self):
        if (REPO / "validator_loader.py").exists():
            from validator_loader import load_validator_core, shared_validator_path
            imported = load_validator_core(shared_validator_path())
            self.assertEqual(imported.VALIDATOR_SHA256, core.VALIDATOR_SHA256)
            self.assertEqual(Path(imported.__file__).resolve(), Path(core.__file__).resolve())
            return
        entrypoint = REPO / ("agent_ref_core.py" if (REPO / "agent_ref_core.py").exists() else "app.py")
        source = entrypoint.read_text(encoding="utf-8").split("# =========================\n# Streamlit UI", 1)[0]
        namespace = {"__file__": str(entrypoint)}
        exec(compile(source, str(entrypoint), "exec"), namespace)
        for name in ("validate_single_ref", "split_references", "score_result"):
            self.assertIs(namespace[name], getattr(core, name))

    def test_exports_accept_shared_results_and_include_release_provenance(self):
        import pandas as pd
        from openpyxl import load_workbook
        result = asyncio.run(core.validate_single_ref(None, "World Health Organization (2024) Diabetes. Available at: https://www.who.int/news-room/fact-sheets/detail/diabetes"))
        self.assertEqual(result["Validation Result"], core.MANUAL_REVIEW_LABEL)
        self.assertEqual(result["Validator Version"], core.VALIDATOR_VERSION)
        result["Ref #"] = 1
        for view in ("Full", "Condensed"):
            with self.subTest(view=view):
                data = build_excel_workbook(pd.DataFrame([result]), view=view)
                workbook = load_workbook(io.BytesIO(data))
                headers = [cell.value for cell in workbook["Results"][1]]
                self.assertIn("Validator Version", headers)
                self.assertIn("Validator SHA256", headers)
                self.assertEqual(workbook["Issues"].max_row, 2)
