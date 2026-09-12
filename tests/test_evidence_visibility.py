"""Uncertain evidence remains visible in interface comparisons and exports."""
import ast
import io
from pathlib import Path
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from agent_ref_validator import core
from agent_ref_validator.exports import build_excel_workbook


def uncertain_result():
    return {"Ref #": 1, "Reference": "Example, A. (2020) ‘Example title’, Journal of Examples.",
        "Validation Result": core.UNVERIFIED_LABEL, "Score": 7, "Domain": "", "AI URL flag": False,
        "Notes": "unusable author metadata", "Extracted Authors": "Example", "Metadata Authors": "Head",
        "author_ok": False, "Author match count": 0, "Submitted author count": 1,
        "Comparison evidence": {"authors": {"status": "unknown", "reason": "unusable author metadata"}},
        "Validator Version": core.VALIDATOR_VERSION, "Validator SHA256": core.VALIDATOR_SHA256}


class EvidenceVisibilityTests(unittest.TestCase):
    def test_unverified_reference_is_in_full_and_condensed_issues_exports(self):
        result = uncertain_result()
        for view in ("Full", "Condensed"):
            with self.subTest(view=view):
                book = load_workbook(io.BytesIO(build_excel_workbook(pd.DataFrame([result]), view=view)))
                self.assertEqual(book["Issues"].max_row, 2)
                headers = [c.value for c in book["Issues"][1]]
                values = [c.value for c in book["Issues"][2]]
                self.assertEqual(values[headers.index("Validation Result")], core.UNVERIFIED_LABEL)
                self.assertIn("unusable author metadata", values[headers.index("Notes")])

    def test_part_comparison_displays_uncertainty_not_mismatch(self):
        path = REPO / ("inspector.py" if (REPO / "inspector.py").exists() else "app.py")
        if not path.exists():
            self.skipTest("Document interface displays results and notes without a part comparison")
        names = {"_display_text", "_part_status", "build_part_by_part_df", "_safe", "compare_parts"}
        tree = ast.parse(path.read_text(encoding="utf-8"))
        subset = ast.Module(body=[node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in names], type_ignores=[])
        import html
        from typing import Any
        namespace = {"pd": pd, "html": html, "core": core, "Any": Any, **vars(core)}
        if path.name == "inspector.py":
            from comparison_helpers import compare_doi_status, compare_year_tokens
            namespace.update(compare_doi_status=compare_doi_status, compare_year_tokens=compare_year_tokens)
        exec(compile(subset, str(path), "exec"), namespace)
        result = uncertain_result()
        if path.name == "inspector.py":
            table = namespace["compare_parts"]({}, {}, result)
            row = table.loc[table["Part"] == "Author list"].iloc[0]
            self.assertEqual(row["Check"], "Not verified")
            self.assertEqual(row["Notes"], "unusable author metadata")
        else:
            table = namespace["build_part_by_part_df"](pd.Series(result))
            row = table.loc[table["Part"] == "Authors"].iloc[0]
            self.assertEqual(row["Status"], "Not verified")
            self.assertEqual(row["What this means"], "unusable author metadata")


if __name__ == "__main__":
    unittest.main()
